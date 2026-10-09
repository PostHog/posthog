import io
import gzip
import hashlib
from collections.abc import Iterable
from dataclasses import replace
from datetime import UTC, datetime
from http import HTTPStatus
from typing import Any, Optional, Union, cast

import pytest
from unittest import mock

import requests
import responses
from openpyxl import Workbook
from parameterized import parameterized
from structlog.testing import capture_logs
from urllib3.response import HTTPResponse

from products.warehouse_sources.backend.models.external_data_schema import SCHEMA_RESOURCE_ID_METADATA_KEY
from products.warehouse_sources.backend.temporal.data_imports.sources.common.excel_parsing import (
    EXCEL_ERROR,
    ExcelFileError,
)
from products.warehouse_sources.backend.temporal.data_imports.sources.common.resumable import ResumableSourceManager
from products.warehouse_sources.backend.temporal.data_imports.sources.common.typings import SourceInputs
from products.warehouse_sources.backend.temporal.data_imports.sources.generated_configs.sharepoint import (
    SharePointSourceConfig,
)
from products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.canonical_descriptions import (
    CANONICAL_DESCRIPTIONS,
)
from products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.files import (
    SharePointFile,
    discover_files,
    files_by_table,
    sharepoint_file_source,
)
from products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.settings import (
    ENDPOINTS,
    FILE_NOT_FOUND_ERROR,
    GRAPH_BASE_URL,
    MAX_FILES,
    PATTERN_ERROR,
    SHAREPOINT_ENDPOINTS,
)
from products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.sharepoint import (
    INVALID_SITE_URL_ERROR,
    SITES_DENIED_ERROR,
    SharePointClient,
    SharePointResumeConfig,
    SharePointSiteURLError,
    _get_rows,
    parse_site_urls,
    validate_credentials,
)
from products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.source import SharePointSource

MODULE = "products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.sharepoint"
FILES_MODULE = "products.warehouse_sources.backend.temporal.data_imports.sources.sharepoint.files"

TENANT_ID = "11111111-1111-1111-1111-111111111111"
CLIENT_ID = "22222222-2222-2222-2222-222222222222"
CLIENT_SECRET = "super-secret"

SITE_A = "contoso.sharepoint.com,aaaa,1111"
SITE_B = "contoso.sharepoint.com,bbbb,2222"

Route = Union[dict[str, Any], int, bytes]


def _response(status: int = 200, json_body: Any = None) -> mock.MagicMock:
    response = mock.MagicMock(spec=requests.Response)
    response.status_code = status
    response.reason = HTTPStatus(status).phrase
    response.ok = 200 <= status < 300
    response.text = ""
    response.json.return_value = json_body if json_body is not None else {}
    if not response.ok:
        response.raise_for_status.side_effect = requests.HTTPError(f"{status} Client Error", response=response)
    else:
        response.raise_for_status.return_value = None
    return response


def _token_response() -> mock.MagicMock:
    return _response(200, {"access_token": "token-1", "expires_in": 3599})


def _session(routes: dict[str, Route], post_responses: Optional[list[mock.MagicMock]] = None) -> Any:
    session = mock.MagicMock(spec=requests.Session)

    def _get(url: str, **kwargs: Any) -> mock.MagicMock:
        route = routes[url]
        if isinstance(route, int):
            return _response(route)
        if isinstance(route, bytes):
            response = _response()
            response.raw = HTTPResponse(body=io.BytesIO(route), preload_content=False)
            return response
        return _response(200, route)

    session.get.side_effect = _get
    session.post.side_effect = post_responses if post_responses is not None else [_token_response()]
    return session


def _manager(resume_state: Optional[SharePointResumeConfig] = None) -> Any:
    manager = mock.MagicMock(spec=ResumableSourceManager)
    manager.can_resume.return_value = resume_state is not None
    manager.load_state.return_value = resume_state
    return manager


def _rows(
    endpoint: str, session: Any, manager: Optional[Any] = None, site_urls: Optional[str] = None
) -> list[dict[str, Any]]:
    with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
        pages = list(
            _get_rows(
                tenant_id=TENANT_ID,
                client_id=CLIENT_ID,
                client_secret=CLIENT_SECRET,
                site_urls=site_urls,
                endpoint=endpoint,
                logger=mock.MagicMock(),
                resumable_source_manager=manager if manager is not None else _manager(),
            )
        )
    return [row for page in pages for row in page]


def _g(path: str) -> str:
    return f"{GRAPH_BASE_URL}{path}"


ALL_SITES_ROUTES: dict[str, Route] = {
    _g("/sites/getAllSites"): {
        "value": [
            {"id": SITE_A, "webUrl": "https://contoso.sharepoint.com/sites/a", "isPersonalSite": False},
            {"id": "personal-1", "webUrl": "https://contoso-my.sharepoint.com/personal/x", "isPersonalSite": True},
        ],
        "@odata.nextLink": _g("/sites/getAllSites?$skiptoken=page2"),
    },
    _g("/sites/getAllSites?$skiptoken=page2"): {
        "value": [
            # Older tenants can omit `isPersonalSite`, so the -my host alone must exclude OneDrive.
            {"id": "personal-2", "webUrl": "https://contoso-my.sharepoint.com/personal/y"},
            {"id": SITE_B, "webUrl": "https://contoso.sharepoint.com/sites/b"},
        ],
    },
}


class TestParseSiteUrls:
    @parameterized.expand(
        [
            ("empty", None, []),
            ("blank_lines", "\n  \n", []),
            ("root_site", "https://contoso.sharepoint.com", ["/sites/contoso.sharepoint.com"]),
            (
                "site_path_with_trailing_slash",
                "https://contoso.sharepoint.com/sites/marketing/",
                ["/sites/contoso.sharepoint.com:/sites/marketing:"],
            ),
            (
                "space_in_path_is_encoded",
                "https://contoso.sharepoint.com/sites/Sales Team",
                ["/sites/contoso.sharepoint.com:/sites/Sales%20Team:"],
            ),
            (
                "newline_and_comma_separated",
                "https://contoso.sharepoint.com/sites/a\nhttps://contoso.sharepoint.com/sites/b, https://contoso.sharepoint.com",
                [
                    "/sites/contoso.sharepoint.com:/sites/a:",
                    "/sites/contoso.sharepoint.com:/sites/b:",
                    "/sites/contoso.sharepoint.com",
                ],
            ),
        ]
    )
    def test_builds_graph_site_paths(self, _name: str, raw: Optional[str], expected: list[str]) -> None:
        assert parse_site_urls(raw) == expected

    @parameterized.expand([("http", "http://contoso.sharepoint.com/sites/a"), ("no_scheme", "contoso.sharepoint.com")])
    def test_rejects_non_https_urls(self, _name: str, raw: str) -> None:
        with pytest.raises(SharePointSiteURLError, match=INVALID_SITE_URL_ERROR):
            parse_site_urls(raw)


class TestSharePointRows:
    @pytest.mark.parametrize("endpoint", ["sites", "drives"])
    def test_explicit_personal_site_is_skipped(self, endpoint: str) -> None:
        site_url = "https://contoso-my.sharepoint.com/personal/user"
        routes: dict[str, Route] = {
            _g("/sites/contoso-my.sharepoint.com:/personal/user:"): {
                "id": "personal-1",
                "webUrl": site_url,
                "isPersonalSite": True,
            }
        }

        assert _rows(endpoint, _session(routes), site_urls=site_url) == []

    def test_next_link_outside_graph_is_refused(self) -> None:
        routes: dict[str, Route] = {
            _g("/sites/getAllSites"): {"value": [{"id": SITE_A}], "@odata.nextLink": "https://evil.example.com/next"}
        }

        with pytest.raises(ValueError, match="non-Graph URL"):
            _rows("sites", _session(routes))

    def test_list_items_skip_hidden_lists_and_carry_parent_ids(self) -> None:
        routes: dict[str, Route] = {
            _g(f"/sites/contoso.sharepoint.com:/sites/a:"): {"id": SITE_A},
            _g(f"/sites/{SITE_A}/lists"): {
                "value": [
                    {"id": "list-1", "list": {"hidden": False}},
                    {"id": "hidden-list", "list": {"hidden": True}},
                    {"id": "deleted-list"},
                ]
            },
            _g(f"/sites/{SITE_A}/lists/list-1/items"): {
                "value": [{"id": "1", "createdDateTime": "2026-01-02T03:04:05Z", "fields": {"Title": "Widget"}}],
            },
            _g(f"/sites/{SITE_A}/lists/deleted-list/items"): 404,
        }
        session = _session(routes)

        rows = _rows("list_items", session, site_urls="https://contoso.sharepoint.com/sites/a")

        assert rows == [
            {
                "id": "1",
                "createdDateTime": datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC),
                "fields": {"Title": "Widget"},
                "site_id": SITE_A,
                "list_id": "list-1",
            }
        ]
        item_calls = [c for c in session.get.call_args_list if c.args[0].endswith("/list-1/items")]
        assert item_calls[0].kwargs["params"] == {"expand": "fields"}
        assert all("hidden-list" not in c.args[0] for c in session.get.call_args_list)

    def test_drive_items_drop_deleted_and_repeated_items(self) -> None:
        routes: dict[str, Route] = {
            **ALL_SITES_ROUTES,
            _g(f"/sites/{SITE_A}/drives"): {"value": [{"id": "b!drive-a"}]},
            _g(f"/sites/{SITE_B}/drives"): {"value": []},
            _g("/drives/b!drive-a/root/delta"): {
                "value": [
                    {"id": "root", "root": {}},
                    {"id": "f1", "file": {}, "@microsoft.graph.downloadUrl": "https://temporary.example.com/file"},
                    {"id": "gone", "deleted": {}},
                ],
                "@odata.nextLink": _g("/drives/b!drive-a/root/delta?token=2"),
            },
            _g("/drives/b!drive-a/root/delta?token=2"): {"value": [{"id": "f1", "file": {}}, {"id": "f2", "file": {}}]},
        }

        rows = _rows("drive_items", _session(routes))

        assert [(row["drive_id"], row["id"]) for row in rows] == [
            ("b!drive-a", "root"),
            ("b!drive-a", "f1"),
            ("b!drive-a", "f2"),
        ]
        assert all(row["site_id"] == SITE_A for row in rows)
        assert all("@microsoft.graph.downloadUrl" not in row for row in rows)

    def test_expired_token_is_reminted_once(self) -> None:
        session = _session({}, post_responses=[_token_response(), _token_response()])
        session.get.side_effect = [_response(401), _response(200, {"value": [{"id": SITE_A}]})]

        rows = _rows("sites", session)

        assert [row["id"] for row in rows] == [SITE_A]
        assert session.post.call_count == 2


class TestSharePointResume:
    def test_resume_on_a_site_that_is_gone_restarts_from_the_first_site(self) -> None:
        routes: dict[str, Route] = {
            **ALL_SITES_ROUTES,
            _g(f"/sites/{SITE_A}/lists"): {"value": [{"id": "list-a"}]},
            _g(f"/sites/{SITE_B}/lists"): {"value": [{"id": "list-b"}]},
        }
        manager = _manager(SharePointResumeConfig(site_id="deleted-site", next_url=_g("/stale")))

        rows = _rows("lists", _session(routes), manager=manager)

        assert [(row["site_id"], row["id"]) for row in rows] == [(SITE_A, "list-a"), (SITE_B, "list-b")]


class TestValidateCredentials:
    @parameterized.expand(
        [
            ("token_rejected", 401, None, None, "Entra ID rejected the app registration"),
            ("token_bad_request", 400, None, None, "Entra ID rejected the app registration"),
            ("graph_forbidden", 200, 403, None, SITES_DENIED_ERROR),
            ("site_not_found", 200, 404, "https://contoso.sharepoint.com/sites/a", "couldn't find the site"),
            ("ok", 200, 200, None, None),
        ]
    )
    def test_maps_status_to_message(
        self,
        _name: str,
        token_status: int,
        graph_status: Optional[int],
        site_urls: Optional[str],
        expected_error: Optional[str],
    ) -> None:
        token = _token_response() if token_status == 200 else _response(token_status)
        session = _session({}, post_responses=[token])
        session.get.side_effect = [_response(graph_status or 200, {"value": []})]

        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            ok, error = validate_credentials(TENANT_ID, CLIENT_ID, CLIENT_SECRET, site_urls)

        assert ok is (expected_error is None)
        if expected_error is None:
            assert error is None
        else:
            assert error is not None and expected_error in error

    def test_invalid_site_url_fails_before_any_request(self) -> None:
        session = _session({})

        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            ok, error = validate_credentials(TENANT_ID, CLIENT_ID, CLIENT_SECRET, "contoso.sharepoint.com/sites/a")

        assert not ok
        assert error is not None and error.startswith(INVALID_SITE_URL_ERROR)
        session.post.assert_not_called()


def test_canonical_descriptions_document_every_primary_key() -> None:
    assert set(CANONICAL_DESCRIPTIONS) == set(ENDPOINTS)
    for name, entry in CANONICAL_DESCRIPTIONS.items():
        assert set(SHAREPOINT_ENDPOINTS[name].primary_keys) <= set(entry["columns"])


SITE_URL = "https://contoso.sharepoint.com/sites/a"


def excel_bytes() -> bytes:
    workbook = Workbook()
    try:
        first = workbook.worksheets[0]
        first.title = "Sales - North"
        first.append(["Order ID", "total"])
        first.append([1, 20])
        workbook.create_sheet("Hidden").sheet_state = "hidden"
        workbook.create_sheet("Sales South").append(["id"])
        workbook.create_sheet("Very hidden").sheet_state = "veryHidden"
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    finally:
        workbook.close()


def file_routes(items: list[dict[str, object]]) -> dict[str, Route]:
    return {
        _g("/sites/contoso.sharepoint.com:/sites/a:"): {"id": SITE_A, "displayName": "Site A"},
        _g(f"/sites/{SITE_A}/drives"): {"value": [{"id": "drive-a", "name": "Shared Documents"}]},
        _g("/drives/drive-a/root/delta"): {"value": items},
    }


def _discover(session: mock.MagicMock, pattern: str | None = None, site_urls: str = SITE_URL) -> list[SharePointFile]:
    with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
        logger = mock.MagicMock()
        client = SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger)
        return discover_files(client, site_urls, pattern, logger)


class TestSharePointFileDiscovery:
    @parameterized.expand(
        [
            ("nested", "nested", "reports/2026/orders.csv"),
            ("root", "root", "orders.csv"),
            ("missing", "missing", "orders.csv"),
            ("missing_ancestor", "orphan", "orders.csv"),
            ("cycle", "cycle", "orders.csv"),
        ]
    )
    def test_resolves_parent_ids_across_pages(self, _name: str, parent: str, expected: str) -> None:
        routes = file_routes([{"id": "file-1", "name": "orders.csv", "file": {}, "parentReference": {"id": parent}}])
        cast(dict[str, object], routes[_g("/drives/drive-a/root/delta")])["@odata.nextLink"] = _g(
            "/drives/drive-a/root/delta?page=2"
        )
        routes[_g("/drives/drive-a/root/delta?page=2")] = {
            "value": [
                {"id": "root", "name": "Ignored root name", "root": {}, "folder": {}},
                {"id": "reports", "name": "reports", "folder": {}, "parentReference": {"id": "root"}},
                {"id": "nested", "name": "2026", "folder": {}, "parentReference": {"id": "reports"}},
                {"id": "orphan", "name": "lost", "folder": {}, "parentReference": {"id": "missing"}},
                {"id": "cycle", "name": "loop", "folder": {}, "parentReference": {"id": "cycle"}},
            ]
        }

        files = _discover(_session(routes))

        assert [file.path for file in files] == [f"Shared Documents/{expected}"]

    @parameterized.expand(
        [
            ("all", None, ["latest.CSV", "compressed.csv.gz", "tabs.TSV.GZ", "book.XLSX"]),
            ("search", r"Documents/latest", ["latest.CSV"]),
            ("no_match", r"^reports/", []),
        ]
    )
    def test_filters_final_delta_state(self, _name: str, pattern: str | None, expected: list[str]) -> None:
        routes = file_routes(
            [
                {"id": "renamed", "name": "old.csv", "file": {}},
                {"id": "deleted-later", "name": "gone.csv", "file": {}},
                {"id": "folder", "name": "folder.csv", "folder": {}},
                {"id": "text", "name": "notes.txt", "file": {}},
                {"id": "json", "name": "data.json", "file": {}},
                {"id": "xls", "name": "legacy.xls", "file": {}},
                {"id": "xlsm", "name": "macros.xlsm", "file": {}},
                {"id": "xlsx-gz", "name": "book.xlsx.gz", "file": {}},
                {"id": "deleted", "name": "deleted.csv", "file": {}, "deleted": {}},
            ]
        )
        cast(dict[str, object], routes[_g("/drives/drive-a/root/delta")])["@odata.nextLink"] = _g(
            "/drives/drive-a/root/delta?page=2"
        )
        routes[_g("/drives/drive-a/root/delta?page=2")] = {
            "value": [
                {"id": "renamed", "name": "latest.CSV", "file": {}},
                {"id": "deleted-later", "deleted": {}},
                {"id": "gzip", "name": "compressed.csv.gz", "file": {}},
                {"id": "tsv", "name": "tabs.TSV.GZ", "file": {}},
                {"id": "xlsx", "name": "book.XLSX", "file": {}},
            ]
        }

        files = _discover(_session(routes), pattern)

        assert [file.path for file in files] == [f"Shared Documents/{name}" for name in expected]

    def test_prefixes_each_site_when_multiple_sites_are_selected(self) -> None:
        routes = file_routes([{"id": "file-a", "name": "orders.csv", "file": {}}])
        routes.update(
            {
                _g("/sites/contoso.sharepoint.com:/sites/b:"): {"id": SITE_B, "name": "Site B"},
                _g(f"/sites/{SITE_B}/drives"): {"value": [{"id": "drive-b", "name": "Shared Documents"}]},
                _g("/drives/drive-b/root/delta"): {"value": [{"id": "file-b", "name": "orders.csv", "file": {}}]},
            }
        )

        files = _discover(_session(routes), site_urls=f"{SITE_URL}\nhttps://contoso.sharepoint.com/sites/b")

        assert [file.path for file in files] == [
            "Site A/Shared Documents/orders.csv",
            "Site B/Shared Documents/orders.csv",
        ]

    def test_stops_after_the_matched_file_limit(self) -> None:
        items: list[dict[str, Any]] = [
            {"id": "unmatched", "name": "skip.csv", "file": {}},
            *[{"id": str(index), "name": f"match-{index}.csv", "file": {}} for index in range(MAX_FILES + 1)],
        ]
        session = _session(file_routes(items))
        logger = mock.MagicMock()
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            client = SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger)
            files = discover_files(client, SITE_URL, "match-", logger)

        assert len(files) == MAX_FILES
        assert files[-1].item_id == str(MAX_FILES - 1)
        logger.warning.assert_called_once()


def _file(path: str, item_id: str = "item-a") -> SharePointFile:
    return SharePointFile(drive_id="drive-a", item_id=item_id, path=path, size=10, modified_at=None)


class TestSharePointFileNames:
    @parameterized.expand(
        [
            ("plain", "Shared Documents/Reports/Orders.CSV.GZ", "shared_documents_reports_orders"),
            ("static_collision", "sites.csv", "sites_bb1b37bb"),
            ("long", "a" * 120 + ".csv", "a" * 100),
        ]
    )
    def test_names_files(self, _name: str, path: str, expected: str) -> None:
        file = _file(path)
        assert files_by_table([file]) == {expected: file}

    @parameterized.expand(
        [
            ("normalized", "Orders 2026.csv", "orders-2026.csv", "orders_2026"),
            ("truncated", "a" * 101 + ".csv", "a" * 102 + ".csv", "a" * 91),
        ]
    )
    def test_suffixes_every_collision_independently_of_order(
        self, _name: str, first_path: str, second_path: str, base: str
    ) -> None:
        files = [_file(first_path, "item-a"), _file(second_path, "item-b")]
        expected = {
            f"{base}_{hashlib.sha256(f'{file.drive_id}:{file.item_id}'.encode()).hexdigest()[:8]}": file
            for file in files
        }

        assert files_by_table(files) == expected
        assert files_by_table(list(reversed(files))) == expected
        assert all(len(name) <= 100 for name in expected)

    def test_literal_hash_suffix_does_not_replace_another_file(self) -> None:
        suffix = hashlib.sha256(b"drive-a:item-a").hexdigest()[:8]
        files = [_file("orders.csv"), _file("orders.tsv", "item-b"), _file(f"orders_{suffix}.csv", "item-c")]

        tables = files_by_table(files)

        assert set(tables.values()) == set(files)
        assert tables == files_by_table(list(reversed(files)))

    @parameterized.expand(
        [
            ("long_file", "a" * 120, "Sales - North", "a" * 88 + "_sales_north"),
            ("long_worksheet", "report", "w" * 120, "r_" + "w" * 98),
        ]
    )
    def test_truncates_file_before_worksheet(self, _name: str, path: str, worksheet: str, expected: str) -> None:
        file = replace(_file(path + ".xlsx"), worksheet=worksheet)

        assert files_by_table([file]) == {expected: file}

    def test_worksheet_collisions_hash_the_full_resource_id(self) -> None:
        files = [replace(_file("a" * 120 + ".xlsx"), worksheet=title) for title in ("Sales - North", "Sales North")]
        expected = {
            f"{'a' * 79}_sales_north_{hashlib.sha256(file.resource_id.encode()).hexdigest()[:8]}": file
            for file in files
        }

        assert files_by_table(files) == expected
        assert files_by_table(list(reversed(files))) == expected
        assert all(len(name) == 100 for name in expected)


class TestSharePointFileSync:
    def test_worksheet_rows_include_file_metadata(self) -> None:
        content = excel_bytes()
        session = _session(
            {
                _g("/drives/drive-a/items/item-a"): {
                    "name": "renamed.xlsx",
                    "size": len(content),
                    "lastModifiedDateTime": "2026-01-02T03:04:05Z",
                },
                _g("/drives/drive-a/items/item-a/content"): content,
            }
        )
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            logger = mock.MagicMock()
            file = replace(_file("old.xlsx"), worksheet="Sales - North")
            response = sharepoint_file_source(
                SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger), "old_table", file.resource_id, logger
            )
            rows = [row for chunk in cast(Iterable[list[dict[str, object]]], response.items()) for row in chunk]

        assert rows == [
            {
                "order_id": 1,
                "total": 20,
                "_file_name": "renamed.xlsx",
                "_file_modified_at": datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC),
            }
        ]

    @parameterized.expand([("reported", 51), ("grew", 1)])
    def test_rejects_oversized_workbook_at_sync(self, _name: str, reported_size: int) -> None:
        session = _session(
            {
                _g("/drives/drive-a/items/item-a"): {"name": "report.xlsx", "size": reported_size},
                _g("/drives/drive-a/items/item-a/content"): excel_bytes(),
            }
        )
        with (
            mock.patch(f"{MODULE}.make_tracked_session", return_value=session),
            mock.patch(f"{FILES_MODULE}.MAX_EXCEL_FILE_BYTES", 50),
        ):
            logger = mock.MagicMock()
            response = sharepoint_file_source(
                SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger),
                "report",
                "drive-a:item-a:Sales - North",
                logger,
            )
            with pytest.raises(ExcelFileError, match=f"^{EXCEL_ERROR}"):
                list(cast(Iterable[object], response.items()))

    @parameterized.expand(
        [
            ("csv", "renamed.csv", b"Order ID,total\n1,20\n"),
            ("tsv", "renamed.tsv", b"Order ID\ttotal\n1\t20\n"),
            ("gzip", "renamed.csv.gz", gzip.compress(b"Order ID,total\n1,20\n")),
        ]
    )
    def test_streams_rows_with_current_file_metadata(self, _name: str, name: str, content: bytes) -> None:
        download = _response()
        download.raw = HTTPResponse(body=io.BytesIO(content), preload_content=False)
        session = _session({})
        session.get.side_effect = [
            _response(200, {"name": name, "size": len(content), "lastModifiedDateTime": "2026-01-02T03:04:05Z"}),
            download,
        ]
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            logger = mock.MagicMock()
            response = sharepoint_file_source(
                SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger),
                "original_table",
                "drive-a:item-a",
                logger,
            )
            rows = [row for chunk in cast(Iterable[list[dict[str, object]]], response.items()) for row in chunk]

        assert rows == [
            {
                "order_id": "1",
                "total": "20",
                "_file_name": name,
                "_file_modified_at": datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC),
            }
        ]
        assert response.name == "original_table"
        assert response.primary_keys is None
        assert response.supports_resume is False
        assert session.get.call_args.args == (_g("/drives/drive-a/items/item-a/content"),)
        assert session.get.call_args.kwargs["stream"] is True
        assert session.get.call_args.kwargs["headers"]["Accept"] == "*/*"
        download.json.assert_not_called()
        download.close.assert_called_once()
        assert download.raw.decode_content is True

    @parameterized.expand(
        [
            ("missing_id", None, 0),
            ("malformed_id", "drive-a:", 0),
            ("deleted", "drive-a:item-a", 1),
            ("deleted_excel", "drive-a:item-a:Sales - North", 1),
            ("empty_worksheet", "drive-a:item-a:", 0),
        ]
    )
    def test_reports_missing_files(self, _name: str, resource_id: str | None, expected_requests: int) -> None:
        session = _session({_g("/drives/drive-a/items/item-a"): 404})
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            logger = mock.MagicMock()
            with pytest.raises(ValueError, match=f"^{FILE_NOT_FOUND_ERROR}"):
                response = sharepoint_file_source(
                    SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger), "orders", resource_id, logger
                )
                list(cast(Iterable[object], response.items()))

        assert session.get.call_count == expected_requests

    @parameterized.expand([("refresh_success", 200), ("refresh_rejected", 401)])
    def test_content_request_refreshes_once(self, _name: str, final_status: int) -> None:
        first = _response(401)
        final = _response(final_status)
        session = _session({}, post_responses=[_token_response(), _token_response()])
        session.get.side_effect = [first, final]
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            client = SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, mock.MagicMock())
            if final_status == 401:
                with pytest.raises(requests.HTTPError):
                    client.open_stream("/drives/drive-a/items/item-a/content")
                final.close.assert_called_once()
            else:
                assert client.open_stream("/drives/drive-a/items/item-a/content") is final

        assert session.post.call_count == 2
        assert session.get.call_count == 2
        first.close.assert_called_once()

    def test_closes_download_when_parsing_fails(self) -> None:
        download = _response()
        download.raw = io.BytesIO(b"invalid gzip")
        session = _session({})
        session.get.side_effect = [_response(200, {"name": "orders.csv.gz"}), download]
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            logger = mock.MagicMock()
            response = sharepoint_file_source(
                SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, logger), "orders", "drive-a:item-a", logger
            )
            with pytest.raises(gzip.BadGzipFile):
                list(cast(Iterable[object], response.items()))

        download.close.assert_called_once()

    def test_refuses_non_graph_initial_download_url(self) -> None:
        session = _session({})
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            client = SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, mock.MagicMock())
            with pytest.raises(ValueError, match="non-Graph URL"):
                client.open_stream("https://example.com/content")

        session.get.assert_not_called()
        session.post.assert_not_called()

    @parameterized.expand([("success", 200), ("denied", 403)])
    def test_redirect_drops_authorization_and_keeps_download_urls_out_of_logs(self, _name: str, status: int) -> None:
        download_url = "https://download.example.com/temporary-file-secret?access=temporary-query-secret"
        client = SharePointClient(TENANT_ID, CLIENT_ID, CLIENT_SECRET, mock.MagicMock())
        content_url = _g("/drives/drive-a/items/item-a/content")
        with responses.RequestsMock() as http, capture_logs() as logs:
            http.post(client.token_url, json={"access_token": "token-1"})
            http.get(content_url, status=302, headers={"Location": download_url})
            http.get(download_url, status=status, body=b"id\n1\n")
            http.get(_g("/sites/site-a"), json={"id": "site-a"})

            if status == 200:
                with client.open_stream(content_url) as response:
                    assert response.content == b"id\n1\n"
            else:
                with pytest.raises(requests.HTTPError) as error:
                    client.open_stream(content_url)
                assert "temporary-file-secret" not in str(error.value)
                assert "temporary-query-secret" not in str(error.value)
            client.get("/sites/site-a")

            assert http.calls[1].request.headers["Authorization"] == "Bearer token-1"
            assert "Authorization" not in http.calls[2].request.headers

        assert "temporary-file-secret" not in str(logs)
        assert "temporary-query-secret" not in str(logs)
        assert any(log.get("url") == "REDACTED" for log in logs)
        assert any(log.get("url") == _g("/sites/site-a") for log in logs)


def make_config(
    enabled: bool | None = True, site_urls: str | None = SITE_URL, pattern: str | None = None
) -> SharePointSourceConfig:
    values: dict[str, object] = {
        "tenant_id": TENANT_ID,
        "client_id": CLIENT_ID,
        "client_secret": CLIENT_SECRET,
        "site_urls": site_urls,
    }
    if enabled is not None:
        values["import_files"] = {"enabled": enabled, "file_pattern": pattern}
    return SharePointSourceConfig.from_dict(values)


def make_inputs(schema_name: str, resource_id: str | None) -> SourceInputs:
    return SourceInputs(
        schema_name=schema_name,
        schema_id="schema-id",
        source_id="source-id",
        team_id=1,
        should_use_incremental_field=False,
        db_incremental_field_last_value=None,
        db_incremental_field_earliest_value=None,
        incremental_field=None,
        incremental_field_type=None,
        job_id="job-id",
        logger=mock.MagicMock(),
        reset_pipeline=False,
        schema_metadata={SCHEMA_RESOURCE_ID_METADATA_KEY: resource_id} if resource_id is not None else None,
    )


class TestSharePointSchemas:
    def test_discovers_one_table_per_visible_worksheet(self) -> None:
        routes = file_routes([{"id": "item-a", "name": "report.xlsx", "file": {}}])
        routes[_g("/drives/drive-a/items/item-a/content")] = excel_bytes()
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=_session(routes)):
            schemas = SharePointSource().get_schemas(make_config(), team_id=1)

        assert [
            (schema.name, schema.label, schema.description, schema.schema_metadata)
            for schema in schemas
            if schema.schema_metadata
        ] == [
            (
                "shared_documents_report_sales_north",
                "Shared Documents/report.xlsx [Sales - North]",
                "Rows of the worksheet Sales - North in the Excel file Shared Documents/report.xlsx",
                {SCHEMA_RESOURCE_ID_METADATA_KEY: "drive-a:item-a:Sales - North"},
            ),
            (
                "shared_documents_report_sales_south",
                "Shared Documents/report.xlsx [Sales South]",
                "Rows of the worksheet Sales South in the Excel file Shared Documents/report.xlsx",
                {SCHEMA_RESOURCE_ID_METADATA_KEY: "drive-a:item-a:Sales South"},
            ),
        ]

    @parameterized.expand([("reported_size",), ("download_size",), ("corrupt",), ("missing",)])
    def test_skips_unreadable_workbooks_and_keeps_sibling_csv(self, case: str) -> None:
        routes = file_routes(
            [
                {"id": "item-a", "name": "report.xlsx", "file": {}, "size": 51 if case == "reported_size" else 1},
                {"id": "csv", "name": "orders.csv", "file": {}},
            ]
        )
        routes[_g("/drives/drive-a/items/item-a/content")] = (
            404 if case == "missing" else b"not a zip" if case == "corrupt" else excel_bytes()
        )
        session = _session(routes)
        with (
            mock.patch(f"{MODULE}.make_tracked_session", return_value=session),
            mock.patch(f"{FILES_MODULE}.MAX_EXCEL_FILE_BYTES", 50),
            capture_logs() as logs,
        ):
            schemas = SharePointSource().get_schemas(make_config(), team_id=1)

        assert [schema.name for schema in schemas if schema.schema_metadata] == ["shared_documents_orders"]
        assert any(
            log.get("log_level") == "warning" and log.get("path") == "Shared Documents/report.xlsx" for log in logs
        )
        if case == "reported_size":
            assert not any(call.args[0].endswith("/content") for call in session.get.call_args_list)

    def test_caps_excel_downloads_in_discovery_order(self) -> None:
        routes = file_routes(
            [
                {"id": "first", "name": "first.xlsx", "file": {}},
                {"id": "second", "name": "second.xlsx", "file": {}},
                {"id": "csv", "name": "orders.csv", "file": {}},
            ]
        )
        routes[_g("/drives/drive-a/items/first/content")] = excel_bytes()
        session = _session(routes)
        with (
            mock.patch(f"{MODULE}.make_tracked_session", return_value=session),
            mock.patch(f"{FILES_MODULE}.MAX_EXCEL_FILES", 1),
            capture_logs() as logs,
        ):
            schemas = SharePointSource().get_schemas(make_config(), team_id=1)

        assert [schema.name for schema in schemas if schema.schema_metadata] == [
            "shared_documents_first_sales_north",
            "shared_documents_first_sales_south",
            "shared_documents_orders",
        ]
        assert any(log.get("log_level") == "warning" and log.get("skipped_files") == 1 for log in logs)
        assert not any("second/content" in call.args[0] for call in session.get.call_args_list)

    @parameterized.expand(
        [
            ("absent", None, None),
            ("off", False, None),
            ("static_names", True, ["drives"]),
            ("empty_names", True, []),
        ]
    )
    def test_static_catalog_needs_no_requests(self, _name: str, enabled: bool | None, names: list[str] | None) -> None:
        session = _session({})
        source = SharePointSource()
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            expected = source.get_schemas(make_config(enabled=None, site_urls=None), team_id=1, names=names)
            schemas = source.get_schemas(make_config(enabled=enabled, site_urls=None), team_id=1, names=names)

        assert schemas == expected
        session.get.assert_not_called()
        session.post.assert_not_called()

    @parameterized.expand(
        [
            ("all", None),
            ("combined_names", ["sites", "shared_documents_orders"]),
        ]
    )
    def test_discovers_file_schemas_with_resource_ids(self, _name: str, names: list[str] | None) -> None:
        session = _session(file_routes([{"id": "item-a", "name": "orders.csv", "file": {}}]))
        source = SharePointSource()
        static = source.get_schemas(make_config(enabled=False), team_id=1, names=names)
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            schemas = source.get_schemas(make_config(), team_id=1, names=names)

        assert schemas[:-1] == static
        file = schemas[-1]
        assert file.name == "shared_documents_orders"
        assert file.label == "Shared Documents/orders.csv"
        assert file.description == "Rows of the CSV file Shared Documents/orders.csv"
        assert file.schema_metadata == {SCHEMA_RESOURCE_ID_METADATA_KEY: "drive-a:item-a"}
        assert not file.supports_append and not file.supports_incremental


class TestSharePointCredentialValidation:
    @parameterized.expand(
        [
            ("missing_sites", None, None, "Enter at least one site URL to import file contents."),
            ("blank_sites", " ,\n ", None, "Enter at least one site URL to import file contents."),
            ("invalid_pattern", SITE_URL, "([", PATTERN_ERROR),
        ]
    )
    def test_rejects_invalid_file_settings_without_http(
        self, _name: str, site_urls: str | None, pattern: str | None, expected: str
    ) -> None:
        session = _session({})
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            valid, error = SharePointSource().validate_credentials(make_config(site_urls=site_urls, pattern=pattern), 1)

        assert valid is False
        assert error is not None and error.startswith(expected)
        session.get.assert_not_called()
        session.post.assert_not_called()

    def test_validates_the_site_without_discovering_files(self) -> None:
        session = _session({_g("/sites/contoso.sharepoint.com:/sites/a:"): {"id": SITE_A}})
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            assert SharePointSource().validate_credentials(make_config(pattern=r"\.csv$"), 1) == (True, None)

        assert session.get.call_count == 1


class TestSharePointPipeline:
    @parameterized.expand([("static", "sites", None), ("file", "old_table_name", "drive-a:item-a")])
    def test_routes_static_tables_and_stored_file_ids(
        self, _name: str, schema_name: str, resource_id: str | None
    ) -> None:
        session = _session(
            {
                _g("/sites/contoso.sharepoint.com:/sites/a:"): {"id": SITE_A},
                _g("/drives/drive-a/items/item-a"): {"name": "new-name.csv"},
                _g("/drives/drive-a/items/item-a/content"): b"id\n1\n",
            }
        )
        manager = _manager()
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            response = SharePointSource().source_for_pipeline(
                make_config(), manager, make_inputs(schema_name, resource_id)
            )
            rows = [row for chunk in cast(Iterable[list[dict[str, object]]], response.items()) for row in chunk]

        if resource_id is None:
            assert rows == [{"id": SITE_A}]
        else:
            assert rows == [{"id": "1", "_file_name": "new-name.csv", "_file_modified_at": None}]
            assert response.supports_resume is False
            manager.save_state.assert_not_called()

    def test_missing_metadata_fails_before_http(self) -> None:
        session = _session({})
        with mock.patch(f"{MODULE}.make_tracked_session", return_value=session):
            with pytest.raises(ValueError, match=f"^{FILE_NOT_FOUND_ERROR}"):
                SharePointSource().source_for_pipeline(make_config(), _manager(), make_inputs("orders", None))

        session.get.assert_not_called()
        session.post.assert_not_called()
