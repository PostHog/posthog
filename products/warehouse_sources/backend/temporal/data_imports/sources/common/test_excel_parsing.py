import io
import re
import zipfile
from datetime import date, datetime, time, timedelta

import pytest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from parameterized import parameterized

from products.warehouse_sources.backend.temporal.data_imports.sources.common.excel_parsing import (
    EXCEL_ERROR,
    ExcelFileError,
    ExcelValue,
    iter_worksheet_rows,
    list_worksheets,
    normalize_chunk,
)

MODULE = "products.warehouse_sources.backend.temporal.data_imports.sources.common.excel_parsing"


def workbook_stream(rows: list[list[object]]) -> io.BytesIO:
    workbook = Workbook(iso_dates=True)
    try:
        ws = workbook.worksheets[0]
        for row in rows:
            ws.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream
    finally:
        workbook.close()


class TestExcelParsing:
    def test_lists_visible_worksheets_in_workbook_order(self) -> None:
        workbook = Workbook()
        try:
            first = workbook.worksheets[0]
            first.title = "Z first"
            first.append([1])
            workbook.create_sheet("hidden").sheet_state = "hidden"
            chart = BarChart()
            chart.add_data(Reference(first, min_col=1, min_row=1, max_row=1))
            workbook.create_chartsheet("Chart").add_chart(chart)
            workbook.create_sheet("very hidden").sheet_state = "veryHidden"
            workbook.create_sheet("A last")
            stream = io.BytesIO()
            workbook.save(stream)
        finally:
            workbook.close()

        assert list_worksheets(stream, "report.xlsx") == ["Z first", "A last"]
        assert not stream.closed

    def test_normalizes_headers_and_row_widths(self) -> None:
        stream = workbook_stream(
            [
                ["Order Total (USD)", None, "Name", "name", 42, None, " "],
                [10, "inside", "a", "b", "c", "ignored"],
                [None, " ", "\t"],
                [20],
                [None, None, None, None, None, "outside"],
            ]
        )

        assert list(iter_worksheet_rows(stream, "report.xlsx", "Sheet")) == [
            [
                {"order_total_usd": 10, "column": "inside", "name": "a", "name_2": "b", "42": "c"},
                {"order_total_usd": 20, "column": None, "name": None, "name_2": None, "42": None},
            ]
        ]

    @parameterized.expand([("empty", []), ("blank_header", [[None, " "], [1, 2]])])
    def test_no_headers_yields_nothing(self, _name: str, rows: list[list[object]]) -> None:
        assert list(iter_worksheet_rows(workbook_stream(rows), "report.xlsx", "Sheet")) == []

    @parameterized.expand(
        [
            ("integer", 7, 7),
            ("float", 1.5, 1.5),
            ("boolean", True, True),
            ("text", "hello", "hello"),
            ("datetime", datetime(2026, 1, 2, 3, 4, 5), datetime(2026, 1, 2, 3, 4, 5)),
            ("date", date(2026, 1, 2), "2026-01-02"),
            ("time", time(3, 4, 5), "03:04:05"),
            ("duration", timedelta(days=2, seconds=3), "2 days, 0:00:03"),
            ("uncached_formula", "=1+1", None),
        ]
    )
    def test_cell_values(self, _name: str, value: object, expected: ExcelValue) -> None:
        stream = workbook_stream([["value", "id"], [value, 1]])

        row = next(iter_worksheet_rows(stream, "report.xlsx", "Sheet"))[0]

        assert row == {"value": expected, "id": 1}
        assert type(row["value"]) is type(expected)

    @parameterized.expand(
        [
            ("number_text", [1, "two", None], ["1", "two", None]),
            ("numeric", [1, 2.5], [1, 2.5]),
            ("boolean_number", [True, 1], ["True", "1"]),
            ("null", [None, 1], [None, 1]),
            ("all_null", [None, None], [None, None]),
            ("date_text", [datetime(2026, 1, 2), "later"], ["2026-01-02 00:00:00", "later"]),
        ]
    )
    def test_mixed_columns(self, _name: str, values: list[ExcelValue], expected: list[ExcelValue]) -> None:
        rows = [{"value": value} for value in values]

        assert normalize_chunk(rows) == [{"value": value} for value in expected]
        assert rows == [{"value": value} for value in values]
        sheet_rows: list[list[object]] = [["value", "id"]]
        for index, value in enumerate(values):
            row: list[object] = [value, index]
            sheet_rows.append(row)
        stream = workbook_stream(sheet_rows)
        parsed = list(iter_worksheet_rows(stream, "report.xlsx", "Sheet"))
        assert [row["value"] for chunk in parsed for row in chunk] == expected

    def test_chunks_and_normalizes_each_chunk_independently(self) -> None:
        stream = workbook_stream([["value"], [1], ["two"], [3], [4.5], [True]])

        assert list(iter_worksheet_rows(stream, "report.xlsx", "Sheet", chunk_size=2)) == [
            [{"value": "1"}, {"value": "two"}],
            [{"value": 3}, {"value": 4.5}],
            [{"value": True}],
        ]

    def test_ignores_incorrect_dimensions(self) -> None:
        original = workbook_stream([["first", "second"], [1, 2], [3]])
        stream = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(stream, "w") as target:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', content)
                target.writestr(member, content)

        assert list(iter_worksheet_rows(stream, "report.xlsx", "Sheet")) == [
            [{"first": 1, "second": 2}, {"first": 3, "second": None}]
        ]

    @parameterized.expand([("not_zip",), ("missing_part",), ("shared_strings",), ("invalid_xml",)])
    def test_invalid_workbooks_have_stable_errors(self, case: str) -> None:
        stream = io.BytesIO(b"not a zip")
        if case != "not_zip":
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                if case == "shared_strings":
                    archive.writestr("xl/sharedStrings.xml", b"x" * 100)
                elif case == "invalid_xml":
                    archive.writestr("[Content_Types].xml", b"not xml")

        with patch(f"{MODULE}.MAX_SHARED_STRINGS_BYTES", 50):
            with pytest.raises(ExcelFileError, match=f"^{EXCEL_ERROR}") as error:
                list_worksheets(stream, "broken.xlsx")
            assert "broken.xlsx" in str(error.value)
            if case != "shared_strings":
                assert "Save it as .xlsx again" in str(error.value)
            with pytest.raises(ExcelFileError, match=f"^{EXCEL_ERROR}"):
                list(iter_worksheet_rows(stream, "broken.xlsx", "Sheet"))

    def test_missing_worksheet_has_refresh_instruction(self) -> None:
        with pytest.raises(ExcelFileError, match=f"^{EXCEL_ERROR}") as error:
            list(iter_worksheet_rows(workbook_stream([["id"], [1]]), "report.xlsx", "Old name"))

        assert str(error.value) == (
            f"{EXCEL_ERROR} 'report.xlsx' has no worksheet named 'Old name'. "
            "It may have been renamed or deleted. Refresh the source's tables."
        )
